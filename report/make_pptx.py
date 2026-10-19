"""실험대 발표 자료 (report/report.pptx) — 웹 뷰어 '실험대' 스크린샷 + 단계 설명.

    ./tools.sh web &                              # 웹 뷰어가 떠 있어야 스크린샷을 찍음
    .venv/bin/python report/make_pptx.py          # 스크린샷 → report/img/bench/*.jpg → report/report.pptx
    .venv/bin/python report/make_pptx.py --no-shots   # 저장된 스크린샷으로 슬라이드만 다시

필요: .venv/bin/pip install python-pptx playwright (브라우저는 설치된 google-chrome 을 씀).
단계 내용은 snss_battery/process.py (RECIPE.md 3절) 그대로.
"""

from __future__ import annotations

import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from pptx import Presentation  # noqa: E402
from pptx.dml.color import RGBColor  # noqa: E402
from pptx.enum.shapes import MSO_SHAPE  # noqa: E402
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN  # noqa: E402
from pptx.oxml.ns import qn  # noqa: E402
from pptx.util import Emu, Inches, Pt  # noqa: E402

from snss_battery.process import HAND, process, summary  # noqa: E402

SHOTS = ROOT / "report" / "img" / "bench"
OUT = ROOT / "report" / "report.pptx"
URL = "http://127.0.0.1:8765/#lab"
FONT = "Malgun Gothic"  # 맑은 고딕 — Windows · Office 기본 한글 글꼴

INK, ICE, PAPER, WHITE = "16324F", "DCEBF7", "F4F7FA", "FFFFFF"
HANDC, MACHC, MUTED, LINE = "E07A1F", "2F6FD6", "5B6B7B", "C9D3DD"
W, H = 13.333, 7.5


# ---------- 스크린샷 ----------
def shoot() -> None:
    """실험대 탭을 단계마다 애니메이션 60 % 지점으로 옮겨 무대(.lab-stage)만 찍음."""
    from PIL import Image
    from playwright.sync_api import sync_playwright

    SHOTS.mkdir(parents=True, exist_ok=True)
    width = "() => parseFloat(document.querySelector('#labBar').style.width) / 100"
    with sync_playwright() as p:
        b = p.chromium.launch(channel="chrome", headless=True)
        pg = b.new_page(viewport={"width": 1600, "height": 900}, device_scale_factor=2)
        pg.goto(URL)
        pg.wait_for_timeout(2500)
        pg.screenshot(path=SHOTS / "full.png")
        n = pg.locator("#labSteps li").count()
        starts = [0.0]
        for _ in range(n - 1):
            pg.click("#labNext")
            pg.wait_for_timeout(100)
            starts.append(pg.evaluate(width))
        starts.append(1.0)
        for k in range(n):
            pg.evaluate("u => window.__lab.seek(u)", starts[k] + 0.6 * (starts[k + 1] - starts[k]))
            pg.wait_for_timeout(700)
            pg.locator(".lab-stage").screenshot(path=SHOTS / f"step{k + 1:02d}.png")
        b.close()
    for f in SHOTS.glob("*.png"):  # 아래 빈 서랍장은 잘라 냄
        im = Image.open(f).convert("RGB")
        if f.stem != "full":
            im = im.crop((0, 0, im.width, int(im.width / 2.54)))
        im.save(f.with_suffix(".jpg"), quality=88)
        f.unlink()


# ---------- 슬라이드 도우미 ----------
def rgb(h: str) -> RGBColor:
    return RGBColor.from_string(h)


def _font(run, size, color, bold=False):
    run.font.size, run.font.bold, run.font.color.rgb = Pt(size), bold, rgb(color)
    run.font.name = FONT
    rpr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):  # 한글은 ea 글꼴로 그려짐
        el = rpr.find(qn(tag))
        if el is None:
            el = rpr.makeelement(qn(tag), {})
            rpr.append(el)
        el.set("typeface", FONT)


def text(slide, x, y, w, h, runs, size=16, color=INK, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
         margin=0.0, space=6, name=None):
    """runs: 문자열, 또는 문단 목록 [[(글, {size,color,bold}), ...], ...]."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        tb.name = name
    tf = tb.text_frame
    tf.word_wrap, tf.vertical_anchor = True, anchor
    for side in ("left", "right", "top", "bottom"):
        setattr(tf, f"margin_{side}", Inches(margin))
    paras = [[(runs, {})]] if isinstance(runs, str) else runs
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(space)
        for t, o in para:
            r = p.add_run()
            r.text = t
            _font(r, o.get("size", size), o.get("color", color), o.get("bold", bold))
    return tb


def box(slide, x, y, w, h, fill, line=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.08, name=None):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = rgb(fill)
    if line:
        s.line.color.rgb, s.line.width = rgb(line), Pt(1)
    else:
        s.line.fill.background()
    s.shadow.inherit = False
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = radius
    if name:
        s.name = name
    return s


def chip(slide, x, y, label, fill, w=None, size=13):
    w = w or 0.32 + 0.17 * len(label)
    s = box(slide, x, y, w, 0.38, fill, radius=0.5)
    tf = s.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = label
    _font(r, size, WHITE, True)
    return w


def picture(slide, path, x, y, w=None, h=None, name=None):
    pic = slide.shapes.add_picture(str(path), Inches(x), Inches(y), Inches(w) if w else None, Inches(h) if h else None)
    pic.line.color.rgb, pic.line.width = rgb(LINE), Pt(0.75)
    if name:
        pic.name = name
    return pic


def bg(slide, color):
    f = slide.background.fill
    f.solid()
    f.fore_color.rgb = rgb(color)


def title(slide, t, sub=None):
    text(slide, 0.6, 0.38, W - 1.2, 0.7, t, size=30, bold=True, name="제목")
    if sub:
        text(slide, 0.6, 1.02, W - 1.2, 0.4, sub, size=14, color=MUTED, name="부제")


def dur(h: float) -> str:
    return f"{h:.1f} h" if h >= 1 else f"{h * 60:.0f} 분"


# ---------- 슬라이드 ----------
def build() -> None:
    steps = process()
    sm = summary(steps)
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(W), Inches(H)
    blank = prs.slide_layouts[6]
    new = lambda: prs.slides.add_slide(blank)  # noqa: E731

    # 1. 표지
    s = new()
    bg(s, INK)
    text(s, 0.7, 0.75, 6.0, 0.4, "SNSS 생리식염수 결정 배터리", size=16, color="9EC3E6", bold=True)
    text(s, 0.7, 1.2, 6.0, 1.9, "통째 결정 실험 —\n실험대에서 보는 19단계", size=36, color=WHITE, bold=True, space=0)
    text(s, 0.7, 3.35, 5.6, 1.2, "약국 식염수병을 그릇 삼아 아래에서 위로 얼려, 얼음 결정 하나하나가 "
         "인공세포가 되게 하는 과정을 웹 뷰어 '실험대' 화면으로 따라갑니다.", size=16, color="D5E3F0")
    text(s, 0.7, 6.55, 6.0, 0.4, "웹 뷰어 http://127.0.0.1:8765/#lab · 2026-10-19", size=12, color="8FA9C2")
    picture(s, SHOTS / "full.jpg", 6.75, 1.2, w=6.0, name="실험대 전체 화면")
    text(s, 6.75, 4.75, 6.0, 0.4, "웹 뷰어 '실험대' 탭 — 지금 화면", size=12, color="8FA9C2")
    s.notes_slide.notes_text_frame.text = (
        "웹 뷰어의 실험대 탭을 그대로 찍은 화면입니다. 왼쪽 무대에서 장비가 움직이고, 오른쪽에는 계기, 실험 노트, 절차가 나옵니다. "
        "이 자료의 단계 화면은 모두 같은 탭에서 각 단계의 60 % 지점을 찍은 것입니다.")

    # 2. 개요 — 숫자
    s = new()
    bg(s, WHITE)
    title(s, "한눈에 보는 실험", "500 mL 식염수병 하나 기준 · 동결 · 해동 시간은 FEniCSx 열전달 계산값")
    stats = [(f"{sm['days']:.1f}일", f"전체 약 {sm['hours']:.0f} h"), (f"{sm['hands_h']:.1f} h", "사람 손이 드는 시간"),
             ("19단계", f"수작업 {sm['hand_steps']} · 기계작업 {sm['machine_steps']}"), ("22 h", "방향성 동결 1회 (3회 반복)")]
    for i, (big, small) in enumerate(stats):
        x = 0.6 + i * 3.08
        box(s, x, 1.75, 2.82, 2.05, PAPER, name=f"숫자 {i + 1}")
        text(s, x + 0.25, 1.95, 2.4, 1.0, big, size=40, bold=True, color=INK if i else MACHC)
        text(s, x + 0.25, 3.0, 2.4, 0.6, small, size=14, color=MUTED)
    pts = [("결정 = 세포.", " 아래에서 위로 천천히 얼리면 얼음 결정 하나가 세포 하나의 틀이 되고, GO · PVA 는 결정 경계로 밀려나 막이 됩니다."),
           ("빈 세포에 에너지.", " 소금도 얼음에서 밀려나, 녹인 뒤 세포 안은 묽고(약 2 mM) 막 사이는 진합니다(약 1.6 M). 이 농도차가 담긴 에너지입니다."),
           ("전기는 아직 안 나옵니다.", " 층 · 극성이 없어 위아래가 대칭이라 윗면–아랫면 전압은 약 2 mV 가 정상입니다.")]
    text(s, 0.6, 4.2, W - 1.2, 2.8, [[(a, {"bold": True}), (b, {})] for a, b in pts], size=16, space=10)

    # 3. 재료표
    s = new()
    bg(s, WHITE)
    title(s, "재료표", "모두 금속 무함유 — 소금의 나트륨 이온은 체액 이온 · 500 mL 병 하나 기준")
    rows = [("재료", "규격", "쓰는 양"),
            ("멸균생리식염수", "0.9 %, 500 mL 나사 마개 병", "1병 (병은 그릇으로도 씀)"),
            ("그래핀 산화물(GO) 분산액", "4 mg/mL, 판 10 µm 이상", "62.5 mL → 병 안 0.5 mg/mL"),
            ("PVA (폴리비닐알코올)", "분자량 85,000–124,000, 가수분해 99 % 이상", "15.00 g → 병 전체 3 wt%"),
            ("증류수", "—", "증발한 만큼 보충"),
            ("접촉액", "30 % 에탄올 (−20 °C 에서 안 얾)", "몇 mL · 전지에 안 들어감")]
    tbl = s.shapes.add_table(len(rows), 3, Inches(0.6), Inches(1.65), Inches(8.1), Inches(4.2)).table
    for j, wd in enumerate((2.75, 3.05, 2.3)):
        tbl.columns[j].width = Inches(wd)
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            c = tbl.cell(i, j)
            c.fill.solid()
            c.fill.fore_color.rgb = rgb(INK if i == 0 else (PAPER if i % 2 else WHITE))
            c.margin_left = c.margin_right = Inches(0.12)
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = c.text_frame
            tf.word_wrap = True
            r = tf.paragraphs[0].add_run()
            r.text = val
            _font(r, 14, WHITE if i == 0 else INK, i == 0 or j == 0)
    box(s, 9.05, 1.65, 3.68, 4.2, ICE, name="나눠 담기")
    text(s, 9.3, 1.85, 3.2, 3.9, [[("나눠 담는 양", {"bold": True, "size": 18})],
                                  [("① 식염수 100 mL 를 덜어 PVA 15 g 을 녹임 (15 wt%)", {})],
                                  [("② 식염수 62.5 mL 를 더 덜어 GO 자리를 비움", {})],
                                  [("③ 남은 337.5 + PVA 용액 100 + GO 62.5", {})],
                                  [("= 500 mL", {"bold": True, "size": 22, "color": MACHC})]], size=14, space=10)
    text(s, 0.6, 6.15, W - 1.2, 0.7, "대조군 A · C 는 같은 배합을 한 번에 만들어 나눠 담고, 대조군 B 는 GO 대신 식염수 62.5 mL 로 따로 만듭니다.",
         size=14, color=MUTED)

    # 4. 전체 흐름
    s = new()
    bg(s, WHITE)
    title(s, "네 묶음으로 보는 19단계", "단계 시간은 '걸리는 시간', 괄호는 그중 손이 드는 시간")
    groups = [("준비", "1–10", range(0, 10), "PVA 를 녹이고 GO 와 섞어 병에 다시 담기"),
              ("동결 · 해동", "11–14", range(10, 14), "냉각판에서 아래 → 위 방향성 동결, 수조 해동, 3회"),
              ("꺼내기 · 측정", "15–17", range(14, 17), "병을 자르고 전압 · 단면 사진 · 조각 고정"),
              ("분석", "18–19", range(17, 19), "동결건조 뒤 SEM-EDS 로 소금 분포")]
    for i, (name, rng, idx, desc) in enumerate(groups):
        x = 0.6 + i * 3.08
        hrs = sum(steps[k].hours for k in idx)
        hands = sum(steps[k].hands_min for k in idx) / 60
        box(s, x, 1.75, 2.82, 4.6, PAPER, name=f"묶음 {i + 1}")
        text(s, x + 0.25, 1.95, 2.4, 0.4, f"단계 {rng}", size=13, color=MUTED, bold=True)
        text(s, x + 0.25, 2.3, 2.4, 0.6, name, size=24, bold=True)
        text(s, x + 0.25, 3.0, 2.4, 0.6, f"{hrs:.1f} h", size=32, bold=True, color=MACHC)
        text(s, x + 0.25, 3.65, 2.4, 0.4, f"(손 {hands:.1f} h)", size=13, color=MUTED)
        text(s, x + 0.25, 4.2, 2.35, 2.0, desc, size=14)
        if i < 3:  # 묶음 사이 화살표
            arr = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x + 2.86), Inches(3.82), Inches(0.18), Inches(0.3))
            arr.fill.solid()
            arr.fill.fore_color.rgb = rgb(MUTED)
            arr.line.fill.background()

    # 5. 단계별
    for k, st in enumerate(steps, 1):
        s = new()
        bg(s, WHITE)
        hand = st.mode == HAND
        text(s, 0.6, 0.3, 9.3, 0.6, [[(f"{k}. ", {"color": MUTED}), (st.title, {})]], size=28, bold=True, name="제목")
        cw = chip(s, W - 0.6 - 1.25, 0.4, "수작업" if hand else "기계작업", HANDC if hand else MACHC, w=1.25)
        text(s, W - 0.6 - cw - 3.1, 0.42, 3.0, 0.38, f"걸림 {dur(st.hours)} · 손 {st.hands_min:.0f} 분", size=14,
             color=MUTED, align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)
        picture(s, SHOTS / f"step{k:02d}.jpg", 0.6, 1.05, w=W - 1.2, name=f"실험대 화면 {k}")
        y = 1.05 + (W - 1.2) / 2.54 + 0.22
        box(s, 0.6, y, 7.45, H - y - 0.35, PAPER, name="할 일")
        text(s, 0.82, y + 0.12, 7.05, H - y - 0.55, [[("할 일  ", {"bold": True, "color": HANDC if hand else MACHC}),
                                                      (st.action, {})]], size=14, anchor=MSO_ANCHOR.MIDDLE)
        box(s, 8.25, y, W - 8.85, H - y - 0.35, ICE, name="넘어가는 조건")
        text(s, 8.47, y + 0.12, W - 9.29, H - y - 0.55, [[("넘어가는 조건  ", {"bold": True, "color": INK}),
                                                          (st.check, {})]], size=14, anchor=MSO_ANCHOR.MIDDLE)
        s.notes_slide.notes_text_frame.text = (f"장비: {' · '.join(st.equipment)}\n할 일: {st.action}\n넘어가는 조건: {st.check}")

    # 마지막 — 판정
    s = new()
    bg(s, INK)
    text(s, 0.7, 0.55, W - 1.4, 0.7, "측정 뒤 무엇을 보나 — 미리 정한 판정 기준", size=30, color=WHITE, bold=True)
    crit = [("E1", "빈 세포에 에너지", "벽 / 칸 Cl 신호 비 10배 이상"), ("E2", "위아래 대칭", "윗면–아랫면 |V| < 50 mV (예측 약 2 mV)"),
            ("S1", "세포와 막", "닫힌 칸, 칸 폭 50–500 µm"), ("S4", "동결이 막을 안 깸", "1회 → 3회 사이 E1 비가 절반 아래로 안 떨어짐")]
    for i, (code, what, rule) in enumerate(crit):
        x, y = 0.7 + (i % 2) * 6.1, 1.65 + (i // 2) * 2.05
        box(s, x, y, 5.85, 1.8, "1F4466", name=f"기준 {code}")
        text(s, x + 0.3, y + 0.22, 1.2, 0.8, code, size=36, color="9EC3E6", bold=True)
        text(s, x + 1.55, y + 0.25, 4.1, 0.5, what, size=20, color=WHITE, bold=True)
        text(s, x + 1.55, y + 0.85, 4.1, 0.8, rule, size=15, color="D5E3F0")
    text(s, 0.7, 6.0, W - 1.4, 1.0, "기준은 측정 전에 정한 것이라 결과를 보고 바꾸지 않습니다. 전체 기준(E1–E3, S1–S4)과 기록 칸은 "
         "인쇄용 기록지 report/report.pdf 와 EXPERIMENT_PROTOCOL.md 에 있습니다.", size=14, color="B9CCDF")

    prs.save(OUT)
    print(f"{OUT} — 슬라이드 {len(prs.slides)}장")


if __name__ == "__main__":
    if "--no-shots" not in sys.argv:
        shoot()
    build()
